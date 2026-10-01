# -*- coding: utf-8 -*-
"""
Automação de Relatório — Solicitações de Compra
Projeto de portfólio — Caíque Augusto Rufino Noel

Lê uma planilha bruta e bagunçada (como normalmente chega de um sistema ou
de preenchimento manual por várias pessoas), trata os dados automaticamente
e gera um relatório final formatado, com resumo e gráficos — tarefa que
manualmente levaria bastante tempo todo mês.
"""

import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

RAW_FILE = "solicitacoes_brutas.xlsx"
OUTPUT_FILE = "Relatorio_Solicitacoes_Compra.xlsx"

log = []  # log de tratamento — vira evidência do que foi corrigido

# =========================================================================
# 1. LER DADO BRUTO
# =========================================================================
df = pd.read_excel(RAW_FILE)
log.append(f"Linhas recebidas no arquivo bruto: {len(df)}")

# =========================================================================
# 2. QUALIDADE DO DADO BRUTO (diagnóstico antes de tratar)
# =========================================================================
n_dup = df.duplicated().sum()
n_missing_valor = df["Valor"].isna().sum()
log.append(f"Linhas duplicadas encontradas: {n_dup}")
log.append(f"Valores faltando na coluna 'Valor': {n_missing_valor}")
log.append(f"Categorias distintas antes da padronização: {df['Categoria'].nunique()}")
log.append(f"Status distintos antes da padronização: {df['Status'].nunique()}")

# =========================================================================
# 3. LIMPEZA
# =========================================================================

# 3.1 Remover duplicatas exatas
df = df.drop_duplicates().reset_index(drop=True)

# 3.2 Padronizar texto (tira espaço extra e uniformiza caixa)
for col in ["Solicitante", "Categoria", "Descrição", "Status"]:
    df[col] = df[col].astype(str).str.strip()

df["Solicitante"] = df["Solicitante"].str.title()
df["Categoria"] = df["Categoria"].str.title()
df["Status"] = df["Status"].str.capitalize()

# normaliza acentuação/variações que "Title Case" sozinho não resolve
categoria_map = {
    "Servicos Graficos": "Serviços Gráficos",
    "Servicos De Ti": "Serviços de TI",
}
df["Categoria"] = df["Categoria"].replace(categoria_map)

# 3.3 Datas em dois formatos diferentes -> padroniza para datetime
def parse_data(valor):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return pd.to_datetime(valor, format=fmt)
        except (ValueError, TypeError):
            continue
    return pd.NaT

df["Data"] = df["Data"].apply(parse_data)

# 3.4 Valor em formatos variados (com "R$", separador de milhar, texto) -> float
def parse_valor(valor):
    if pd.isna(valor):
        return None
    s = str(valor).strip()
    s = re.sub(r"[Rr]\$\s*", "", s)
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None

df["Valor"] = df["Valor"].apply(parse_valor)

# 3.5 Valores faltando: não descarta a solicitação, sinaliza para revisão
df["Valor_Pendente_Revisao"] = df["Valor"].isna()
valores_pendentes = df["Valor_Pendente_Revisao"].sum()
df["Valor"] = df["Valor"].fillna(0)

log.append(f"Linhas duplicadas removidas: {n_dup}")
log.append(f"Categorias distintas após padronização: {df['Categoria'].nunique()}")
log.append(f"Status distintos após padronização: {df['Status'].nunique()}")
log.append(f"Valores sinalizados para revisão manual (não descartados): {valores_pendentes}")
log.append(f"Linhas finais no relatório tratado: {len(df)}")

df = df.sort_values("Data").reset_index(drop=True)
df.to_csv("solicitacoes_tratadas.csv", index=False, sep=";")

print("\n".join(log))

# =========================================================================
# 4. AGREGAÇÕES PARA O RELATÓRIO
# =========================================================================
por_categoria = df.groupby("Categoria")["Valor"].agg(["sum", "count"]).reset_index()
por_categoria.columns = ["Categoria", "Total", "Qtd"]
por_categoria = por_categoria.sort_values("Total", ascending=False)

por_status = df.groupby("Status")["Valor"].agg(["sum", "count"]).reset_index()
por_status.columns = ["Status", "Total", "Qtd"]

df["Mês"] = df["Data"].dt.to_period("M").astype(str)
por_mes = df.groupby("Mês")["Valor"].sum().reset_index()

# =========================================================================
# 5. GERAR RELATÓRIO FORMATADO (.xlsx)
# =========================================================================
NAVY = "1A3D5C"
WHITE = "FFFFFF"
GREY = "F2F2F2"
LIGHT_BLUE = "DCE6F1"
FONT_NAME = "Arial"

header_font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
header_fill = PatternFill("solid", fgColor=NAVY)
normal_font = Font(name=FONT_NAME, size=10)
alt_fill = PatternFill("solid", fgColor=GREY)
kpi_fill = PatternFill("solid", fgColor=LIGHT_BLUE)
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ---- Sheet 1: Sobre / Log de Tratamento ----
ws0 = wb.active
ws0.title = "Sobre e Log"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["B"].width = 90
ws0["B2"] = "Relatório Automatizado de Solicitações de Compra"
ws0["B2"].font = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
ws0["B4"] = "Projeto de portfólio — Caíque Augusto Rufino Noel"
ws0["B4"].font = Font(name=FONT_NAME, size=11, italic=True)
ws0["B6"] = ("Este relatório foi gerado automaticamente por um script Python a partir de um "
             "arquivo bruto e não padronizado (datas em formatos diferentes, valores em texto, "
             "categorias e status escritos de formas diferentes, linhas duplicadas).")
ws0["B6"].alignment = Alignment(wrap_text=True, vertical="top")
ws0.row_dimensions[6].height = 45

ws0["B8"] = "Log de tratamento aplicado:"
ws0["B8"].font = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
for i, linha in enumerate(log):
    ws0.cell(row=9 + i, column=2, value=f"• {linha}").font = normal_font

# ---- Sheet 2: Dados Tratados ----
ws1 = wb.create_sheet("Dados Tratados")
cols = ["Data", "Solicitante", "Categoria", "Descrição", "Valor", "Status", "Valor_Pendente_Revisao"]
for c, h in enumerate(cols, start=1):
    cell = ws1.cell(row=1, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

for r, row in enumerate(df[cols].itertuples(index=False), start=2):
    for c, val in enumerate(row, start=1):
        cell = ws1.cell(row=r, column=c, value=val)
        cell.border = border
        cell.font = normal_font
        if c == 1 and val is not None:
            cell.number_format = "DD/MM/YYYY"
        if c == 5:
            cell.number_format = '"R$" #,##0.00'
        if r % 2 == 0:
            cell.fill = alt_fill

for i, w in enumerate([12, 18, 20, 26, 12, 12, 20], start=1):
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = f"A1:G{len(df)+1}"

# ---- Sheet 3: Resumo ----
ws2 = wb.create_sheet("Resumo")
ws2.sheet_view.showGridLines = False
ws2.merge_cells("B2:E2")
ws2["B2"] = "RESUMO DE SOLICITAÇÕES DE COMPRA"
ws2["B2"].font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
ws2["B2"].fill = PatternFill("solid", fgColor=NAVY)
ws2["B2"].alignment = center
ws2.row_dimensions[2].height = 26

kpis = [("Total Geral", f"R$ {df['Valor'].sum():,.2f}"), ("Qtd. Solicitações", len(df)),
        ("Ticket Médio", f"R$ {df['Valor'].mean():,.2f}"), ("Pendentes de Revisão", int(valores_pendentes))]
for i, (label, value) in enumerate(kpis):
    col = 2 + i
    ws2.cell(row=4, column=col, value=label).font = Font(name=FONT_NAME, size=9, color="444444")
    ws2.cell(row=4, column=col).fill = kpi_fill
    ws2.cell(row=4, column=col).alignment = center
    ws2.cell(row=5, column=col, value=value).font = Font(name=FONT_NAME, size=13, bold=True, color=NAVY)
    ws2.cell(row=5, column=col).fill = kpi_fill
    ws2.cell(row=5, column=col).alignment = center
    ws2.column_dimensions[get_column_letter(col)].width = 18

# Tabela por categoria
ws2["B8"] = "POR CATEGORIA"
ws2["B8"].font = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
for c, h in enumerate(["Categoria", "Total", "Qtd"], start=2):
    cell = ws2.cell(row=9, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = border
    cell.alignment = center
for i, row in enumerate(por_categoria.itertuples(index=False), start=10):
    ws2.cell(row=i, column=2, value=row.Categoria).border = border
    ws2.cell(row=i, column=3, value=row.Total).number_format = '"R$" #,##0.00'
    ws2.cell(row=i, column=3).border = border
    ws2.cell(row=i, column=4, value=row.Qtd).border = border
last_cat_row = 9 + len(por_categoria)

# Gráfico de barras por categoria
bar = BarChart()
bar.title = "Total Gasto por Categoria"
bar.width, bar.height = 16, 9
data = Reference(ws2, min_col=3, min_row=9, max_row=last_cat_row)
cats = Reference(ws2, min_col=2, min_row=10, max_row=last_cat_row)
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
ws2.add_chart(bar, f"F4")

# Gráfico de pizza por status
ws2.cell(row=last_cat_row + 3, column=2, value="POR STATUS").font = Font(
    name=FONT_NAME, size=11, bold=True, color=NAVY)
status_hr = last_cat_row + 4
for c, h in enumerate(["Status", "Total", "Qtd"], start=2):
    cell = ws2.cell(row=status_hr, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = border
    cell.alignment = center
for i, row in enumerate(por_status.itertuples(index=False), start=status_hr + 1):
    ws2.cell(row=i, column=2, value=row.Status).border = border
    ws2.cell(row=i, column=3, value=row.Total).number_format = '"R$" #,##0.00'
    ws2.cell(row=i, column=3).border = border
    ws2.cell(row=i, column=4, value=row.Qtd).border = border
last_status_row = status_hr + len(por_status)

pie = PieChart()
pie.title = "Distribuição por Status"
pie.width, pie.height = 16, 9
data2 = Reference(ws2, min_col=4, min_row=status_hr, max_row=last_status_row)
cats2 = Reference(ws2, min_col=2, min_row=status_hr+1, max_row=last_status_row)
pie.add_data(data2, titles_from_data=True)
pie.set_categories(cats2)
ws2.add_chart(pie, f"F22")

wb.save(OUTPUT_FILE)
print(f"\nRelatório final salvo em: {OUTPUT_FILE}")
